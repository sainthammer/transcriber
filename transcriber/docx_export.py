"""Word export: turns the summary Markdown from `render_markdown` into a styled .docx.

Only the Markdown subset `render_markdown` produces is supported: `#`–`###` headings,
`- ` bullets, pipe tables, `---`, paragraphs with **bold**, _italic_ and `code`
(timestamps). Converting the Markdown rather than the summary model keeps one source
of section order and wording for every summary type.
"""

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Present on both macOS (Pages, Word) and Windows; Calibri isn't installed on a Mac without Office.
FONT = "Arial"
TEXT = RGBColor(0x1D, 0x1D, 0x1F)  # headings and body copy: hierarchy comes from size/weight,
MUTED = RGBColor(0x6E, 0x6E, 0x73)  # not from a second hue (5.1:1 on white, 4.6:1 on TINT)

# One accent family — a light, cool blue in three strengths, never a second hue:
# - SKY is the signature light blue. At 3.3:1 on white it passes WCAG AA for graphics
#   (3:1), not for text, so it only draws: the title rule, table header rule, bullets.
#   It prints as a mid gray (~51% luminance), so rules and bullets survive b/w print.
# - INK is the same blue deepened for text: 6.6:1 on white, 6.0:1 on TINT — small caps
#   labels (meta line, table headers), timestamps and the page number.
# - TINT is a near-white wash of the same hue for the two blocks that earn a surface:
#   the "Кратко" lead and table header rows. It fades to paper in b/w, and nothing
#   depends on it — the text on it is full-contrast ink.
SKY_HEX = "4A90D9"
INK = RGBColor(0x1F, 0x5F, 0x9E)
TINT_HEX = "EEF5FC"
DIVIDER = "DDE3EA"  # cool-gray hairlines: table row lines, the section break

# (size pt, space before, space after) for `#`, `##`, `###`
HEADINGS = {1: (22, 0, 10), 2: (13.5, 24, 7), 3: (11, 14, 4)}
BODY_SIZE = 10.5
# Italic content can't start or end with "_", so model text like `__init__` stays literal.
INLINE = re.compile(r"\*\*(.+?)\*\*|(?<!\w)_([^_\s](?:.*?[^_\s])?)_(?!\w)|`([^`]+)`")
TABLE_SEPARATOR = re.compile(r"^\|(\s*:?-+:?\s*\|)+$")
# ~80 characters of 10.5pt Arial per line: generous but not greedy margins.
MARGIN = Cm(2.4)
CONTENT_WIDTH = Cm(21) - 2 * MARGIN
# Column shares by column count: figures, competencies, action items (# is narrow).
COLUMN_SHARES = {2: (0.58, 0.42), 3: (0.28, 0.2, 0.52), 4: (0.06, 0.5, 0.22, 0.22)}
# Paragraph alignment by column count; only the action items' "#" column isn't left-aligned.
_LEFT, _CENTER = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
COLUMN_ALIGN = {4: (_CENTER, _LEFT, _LEFT, _LEFT)}
# Tables up to this many body rows are kept on one page; longer ones break between rows,
# but never right after the header or before a lone last row. Keeping bigger tables whole
# pushed them to the next page and left half-empty pages behind.
KEEP_TABLE_ROWS = 3
EMPTY_CELL = "—"  # the render's placeholder for "not stated": shown muted, not as data
FOOTER_TITLE_MAX = 80
# Bullet lists up to this long are kept on one page, so a short list under a `###` never
# leaves its last item alone at the top of the next page.
KEEP_LIST_ITEMS = 5


def _set_font(target, name: str) -> None:
    """Set the font for all scripts; theme font attributes would override the name."""
    rpr = target.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), name)
        fonts.attrib.pop(qn(f"w:{attr}Theme"), None)


# w:rPr child order per the OOXML schema, from `w:spacing` onward — python-docx has no
# wrapper for `w:spacing`, so inserting it in the right slot (rather than appending
# blindly, which could land it after `w:sz`/`w:color`) is on us.
_RPR_AFTER_SPACING = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
                      "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
                      "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _set_tracking(target, twips: int) -> None:
    """Character spacing (letter-tracking) in twentieths of a point; negative tightens,
    positive loosens (works on both a style and a run — both expose `.element`)."""
    rpr = target.element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twips))
    for tag in _RPR_AFTER_SPACING:
        successor = rpr.find(qn(tag))
        if successor is not None:
            successor.addprevious(spacing)
            return
    rpr.append(spacing)


def _border(side: str, color: str = DIVIDER, size: int = 4, space: int = 1) -> OxmlElement:
    """size is in eighths of a point, space in points."""
    border = OxmlElement(f"w:{side}")
    for key, value in {"val": "single", "sz": str(size), "space": str(space),
                       "color": color}.items():
        border.set(qn(f"w:{key}"), value)
    return border


def _shading(fill: str) -> OxmlElement:
    shd = OxmlElement("w:shd")
    for key, value in {"val": "clear", "color": "auto", "fill": fill}.items():
        shd.set(qn(f"w:{key}"), value)
    return shd


# w:pPr children that must follow `w:pBdr`/`w:shd` per the schema (python-docx appends
# spacing/ind/jc on demand, so a border added later has to be slotted before them).
_PPR_AFTER_BORDERS = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                      "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                      "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                      "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                      "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                      "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                      "w:pPrChange")


def _insert_ppr(ppr, element, successors: tuple[str, ...]) -> None:
    for tag in successors:
        successor = ppr.find(qn(tag))
        if successor is not None:
            successor.addprevious(element)
            return
    ppr.append(element)


def _paragraph_borders(ppr, *borders) -> None:
    element = OxmlElement("w:pBdr")
    for border in borders:
        element.append(border)
    _insert_ppr(ppr, element, _PPR_AFTER_BORDERS)


def _setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    _set_font(normal, FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True  # no single orphan/widow lines in print
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    normal.element.get_or_add_rPr().append(lang)  # Russian spell check

    for level, (size, before, after) in HEADINGS.items():
        style = doc.styles[f"Heading {level}"]
        _set_font(style, FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = TEXT  # one neutral ink for all levels: size, weight and
        fmt = style.paragraph_format  # air carry the hierarchy — no rule under every heading
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        fmt.keep_with_next = fmt.keep_together = True
        fmt.line_spacing = 1.12 if level == 1 else 1.15  # multi-line titles stay a block
    _set_tracking(doc.styles["Heading 1"], -10)  # tight, confident title tracking
    _set_tracking(doc.styles["Heading 2"], -4)

    bullet = doc.styles["List Bullet"]
    bullet.paragraph_format.space_after = Pt(3)
    # Bullets live in the numbering definition, not the style: a plain Arial "•" (the
    # template's Symbol-font glyph varies by platform) in the light accent, on a tighter
    # 0.5 cm hang than Word's default 0.63 cm so lists sit closer to the text column.
    for lvl in doc.part.numbering_part.element.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        if fmt is None or fmt.get(qn("w:val")) != "bullet" or lvl.get(qn("w:ilvl")) != "0":
            continue
        lvl.find(qn("w:lvlText")).set(qn("w:val"), "•")
        ppr = lvl.find(qn("w:pPr"))
        if ppr is not None:
            for tab in ppr.iter(qn("w:tab")):
                tab.set(qn("w:pos"), "284")
            ind = ppr.find(qn("w:ind"))
            if ind is not None:
                ind.set(qn("w:left"), "284")
                ind.set(qn("w:hanging"), "284")
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            lvl.append(rpr)
        for child in list(rpr):
            rpr.remove(child)
        fonts = OxmlElement("w:rFonts")
        for attr in ("ascii", "hAnsi", "cs"):
            fonts.set(qn(f"w:{attr}"), FONT)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), SKY_HEX)
        rpr.append(fonts)
        rpr.append(color)


def _add_field(paragraph, instruction: str) -> list:
    """A complex field with an explicit result run: apps show the result with that run's
    formatting until they update the field, so every part is returned for styling."""
    runs = []
    for tag, text in (("begin", None), (None, instruction), ("separate", None),
                      ("result", "1"), ("end", None)):
        run = paragraph.add_run()
        if tag == "result":
            run.text = text
        elif tag:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), tag)
            run._r.append(field)
        else:
            field = OxmlElement("w:instrText")
            field.set(qn("xml:space"), "preserve")
            field.text = text
            run._r.append(field)
        runs.append(run)
    return runs


def _setup_page(doc) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.3)
    section.left_margin = section.right_margin = MARGIN
    section.footer_distance = Cm(1.1)


def _setup_footer(doc, title: str) -> None:
    """Title on the left (a loose printed page still says what it belongs to), page
    number on the right. No rule: whitespace already separates it from the text."""
    style = doc.styles["Footer"]
    style.font.size = Pt(8)
    # The template's Footer style has center/right tabs sized for Letter; drop them so the
    # page number lands on our own right tab at the text column's edge.
    tabs = style.element.pPr.find(qn("w:tabs")) if style.element.pPr is not None else None
    if tabs is not None:
        tabs.getparent().remove(tabs)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    footer.paragraph_format.space_after = Pt(0)
    if len(title) > FOOTER_TITLE_MAX:
        title = title[:FOOTER_TITLE_MAX - 1].rstrip() + "…"
    muted = [footer.add_run(f"{title}\t")]
    for run in _add_field(footer, "PAGE"):
        run.bold, run.font.color.rgb = True, INK
    muted += [footer.add_run(" / "), *_add_field(footer, "NUMPAGES")]
    for run in footer.runs:
        run.font.size = Pt(8)
    for run in muted:
        run.font.color.rgb = MUTED


def _add_inline(paragraph, text: str, bold: bool = False) -> None:
    """Add text with **bold**, _italic_ (muted) and `code` (timestamps, small and blue)."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()]).bold = bold or None
        strong, em, code = m.groups()
        if strong is not None:
            paragraph.add_run(strong).bold = True
        elif em is not None:
            run = paragraph.add_run(em)
            run.italic, run.bold = True, bold or None
            run.font.color.rgb = MUTED
        else:
            run = paragraph.add_run(code)
            run.font.size, run.font.color.rgb, run.bold = Pt(8.5), INK, True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:]).bold = bold or None


def _label(run, size: float = 8) -> None:
    """Small tracked capitals in the text accent: meta labels and table headers."""
    run.font.size, run.font.color.rgb, run.font.all_caps, run.bold = Pt(size), INK, True, True
    _set_tracking(run, 14)


def _style_meta(paragraph) -> None:
    """The line under the title: tracked labels, gray values, and the document's one
    strong accent stroke — a light-blue rule that closes the title block."""
    for run in paragraph.runs:
        if run.bold:
            _label(run, 8)
        else:
            run.font.size, run.font.color.rgb = Pt(9.5), MUTED
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(2), Pt(22)
    _paragraph_borders(paragraph._p.get_or_add_pPr(), _border("bottom", SKY_HEX, size=16, space=10))


def _style_lead(paragraph) -> None:
    """The "Кратко" block: a slightly larger lead on a pale tint of the accent. Borders
    in the fill color pad the shading; the indents pull the text in by the same amount,
    so the panel's edges line up with the text column."""
    for run in paragraph.runs:
        run.font.size = Pt(11)
    fmt = paragraph.paragraph_format
    fmt.left_indent = fmt.right_indent = Pt(11)
    fmt.space_before, fmt.space_after = Pt(0), Pt(6)
    fmt.line_spacing = 1.3
    ppr = paragraph._p.get_or_add_pPr()
    _paragraph_borders(ppr, _border("top", TINT_HEX, size=4, space=9),
                       _border("left", TINT_HEX, size=4, space=10),
                       _border("bottom", TINT_HEX, size=4, space=9),
                       _border("right", TINT_HEX, size=4, space=10))
    _insert_ppr(ppr, _shading(TINT_HEX), _PPR_AFTER_BORDERS[1:])


def _cells(row: str) -> list[str]:
    parts = re.split(r"(?<!\\)\|", row.strip())[1:-1]
    return [p.strip().replace("\\|", "|") for p in parts]


def _add_table(doc, rows: list[str], small: bool = False) -> None:
    header, body = _cells(rows[0]), [_cells(r) for r in rows[1:] if not TABLE_SEPARATOR.match(r)]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    # No verticals, no grid: a tinted header band with a light-blue rule under it, neutral
    # hairlines between rows and a closing one at the bottom.
    borders = OxmlElement("w:tblBorders")
    for side in ("bottom", "insideH"):
        borders.append(_border(side))
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, width in (("left", "113"), ("right", "113"), ("top", "85"), ("bottom", "85")):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), width)  # twips
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table._tbl.tblPr.append(margins)

    # Full content width with fixed columns; widths go both to the grid (Pages,
    # Quick Look read it) and to every cell (Word reads those).
    table.autofit = False
    shares = COLUMN_SHARES.get(len(header), [1 / len(header)] * len(header))
    for grid_col, column, share in zip(table._tbl.tblGrid.gridCol_lst, table.columns, shares,
                                       strict=False):
        grid_col.w = int(CONTENT_WIDTH * share)
        for cell in column.cells:
            cell.width = int(CONTENT_WIDTH * share)

    aligns = COLUMN_ALIGN.get(len(header), [_LEFT] * len(header))
    # Keep-with-next on a row's paragraphs binds it to the next row: short tables stay
    # whole; long ones keep the header with two rows and the last two rows together.
    if len(body) <= KEEP_TABLE_ROWS:
        kept = set(range(len(body)))
    else:
        kept = {0, 1, len(body) - 1}
    for r, values in enumerate([header, *body]):
        row = table.rows[r]
        row_pr = row._tr.get_or_add_trPr()
        row_pr.append(OxmlElement("w:cantSplit"))  # a row never breaks across pages
        if r == 0:
            row_pr.append(OxmlElement("w:tblHeader"))  # repeat on every page
        for c, value in enumerate(values[:len(header)]):
            cell = row.cells[c]
            paragraph = cell.paragraphs[0]
            fmt = paragraph.paragraph_format
            fmt.alignment = aligns[c]
            fmt.space_before = fmt.space_after = Pt(0)
            fmt.line_spacing = 1.2
            fmt.keep_with_next = r in kept
            _add_inline(paragraph, value, bold=r == 0)
            if r == 0:  # small tracked capitals on a tint band, light-blue rule below
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_borders = OxmlElement("w:tcBorders")
                tc_borders.append(_border("bottom", SKY_HEX, size=8))
                tc_pr.append(tc_borders)
                tc_pr.append(_shading(TINT_HEX))
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER  # vAlign goes after shd
                for run in paragraph.runs:
                    _label(run, 7.5)
                continue
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for run in paragraph.runs:
                run.font.size = Pt(9.5 if small else 10)
                if value == EMPTY_CELL or (len(header) == 4 and c == 0):
                    run.font.color.rgb = MUTED  # placeholders and row numbers step back
                elif len(header) == 2 and c == 1:
                    run.bold = True  # the figure itself is what the key-figures table is for
    # Schema order: tblLook closes tblPr, after the borders/margins appended above.
    look = table._tbl.tblPr.find(qn("w:tblLook"))
    if look is not None:
        table._tbl.tblPr.append(look)
    # A spacer: Word needs a paragraph after a table, and a bare one would add a full
    # empty line — this one only adds a hair of air before whatever comes next.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(4)


def _add_rule(doc) -> None:
    """A quiet chapter break before "Ход обсуждения": a neutral hairline and extra air —
    what follows is the reference tail, set a notch smaller than the main part."""
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(12), Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(2)
    fmt.keep_with_next = True  # never a lone rule at the bottom of a page
    _paragraph_borders(paragraph._p.get_or_add_pPr(), _border("bottom", DIVIDER, size=6))


def _keep_short_lists(doc) -> None:
    """Bind the items of each short bullet list together (the heading above already
    keeps with the first item), so a list moves to the next page as a whole."""
    items: list = []
    for paragraph in [*doc.paragraphs, None]:
        if paragraph is not None and paragraph.style.name == "List Bullet":
            items.append(paragraph)
            continue
        if 1 < len(items) <= KEEP_LIST_ITEMS:
            for item in items[:-1]:
                item.paragraph_format.keep_with_next = True
        items = []


def markdown_to_docx(md: str) -> bytes:
    doc = Document()
    _setup_styles(doc)
    _setup_page(doc)

    lines = md.splitlines()
    first_paragraph = True  # the line under the title is the meta line
    section = 0  # index of the current `##` section; the first one is the lead ("Кратко")
    tail = False  # after `---`: the discussion log, set a notch smaller
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if heading := re.match(r"^(#{1,3}) (.+)$", line):
            level = len(heading[1])
            if level == 1:
                doc.core_properties.title = heading[2]
            if level == 2:
                section += 1
            paragraph = doc.add_heading(level=level)
            _add_inline(paragraph, heading[2])
            if level == 2 and tail:
                paragraph.paragraph_format.space_before = Pt(14)
            if level == 3 and tail:
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_inline(paragraph, line[2:])
            if tail:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            _add_table(doc, rows, small=tail)
            continue
        elif line.strip() == "---":
            _add_rule(doc)
            tail = True
        else:
            paragraph = doc.add_paragraph()
            _add_inline(paragraph, line)
            if first_paragraph and doc.core_properties.title and section == 0:
                _style_meta(paragraph)
            elif section == 1 and not tail:
                _style_lead(paragraph)
            first_paragraph = False
        i += 1

    _keep_short_lists(doc)
    _setup_footer(doc, doc.core_properties.title or "")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def save_docx(md: str, path: Path) -> None:
    path.write_bytes(markdown_to_docx(md))
