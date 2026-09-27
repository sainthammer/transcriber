from io import BytesIO

import pytest
from docx import Document

from transcriber.docx_export import markdown_to_docx
from transcriber.summary_types import TYPES, render_markdown

from test_summary_types import make_summary


def as_document(md: str):
    return Document(BytesIO(markdown_to_docx(md)))


@pytest.mark.parametrize("summary_type", list(TYPES))
def test_every_type_exports(summary_type):
    md = render_markdown(make_summary(summary_type), duration="~35 мин")
    doc = as_document(md)

    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings[0] == "Пилот ИИ-аналитики"
    assert headings == [line.lstrip("# ").replace("`", "") for line in md.splitlines()
                        if line.startswith("#")]
    assert doc.core_properties.title == "Пилот ИИ-аналитики"

    text = "\n".join(p.text for p in doc.paragraphs)
    for markup in ("**", "`", "| ", "_Не зафиксировано_"):
        assert markup not in text


def test_tables_and_inline_formatting():
    doc = as_document(render_markdown(make_summary("client"), duration="~35 мин"))

    figures, tasks = doc.tables
    assert [c.text for c in figures.rows[1].cells] == ["сервер | оборудование", "1,5 млн руб."]
    assert [c.text for c in tasks.rows[0].cells] == ["#", "Задача", "Ответственный", "Срок"]
    assert [c.text for c in tasks.rows[1].cells] == ["1", "Подготовить вопросы 00:09:39", "—", "среда"]

    meta = doc.paragraphs[1]
    assert meta.text == "Тип: Встреча с клиентом · Длительность: ~35 мин"
    assert meta.runs[0].bold

    participant = next(p for p in doc.paragraphs if p.text.startswith("Максим"))
    assert participant.style.name == "List Bullet"
    assert participant.runs[0].bold and participant.runs[-1].italic  # name, side


def test_underscores_inside_words_are_not_italic():
    doc = as_document("# T\n\n- поле snake_case_name и _курсив_\n")
    item = doc.paragraphs[-1]
    assert item.text == "поле snake_case_name и курсив"
    assert [r.italic for r in item.runs] == [None, True]


def test_freeform_sections_export():
    """Free-form sections (model-chosen titles, paragraphs, points) survive the Word
    download path: timestamps off, empty sections stripped."""
    from transcriber.summary_types import strip_empty_sections

    s = make_summary("freeform", details={"sections": [
        {"title": "Зачем писать постмортем",
         "paragraphs": ["Выводы по инцидентам забываются.", "Документ пишется для дежурных."],
         "points": []},
        {"title": "Структура", "paragraphs": ["Пять частей."],
         "points": ["Хронология", "Причины"]},
        {"title": "Правила", "paragraphs": [], "points": ["Владелец и срок у каждой задачи"]},
        {"title": "Код", "paragraphs": ["Модуль __init__.py и snake_case_name"], "points": []},
    ]}, decisions=[], open_questions=[])
    md = strip_empty_sections(render_markdown(s, duration="~35 мин", timestamps=False))
    doc = as_document(md)
    paragraphs = [(p.style.name, p.text) for p in doc.paragraphs if p.text]

    def body(title):
        start = paragraphs.index(("Heading 2", title)) + 1
        end = next(i for i in range(start, len(paragraphs)) if paragraphs[i][0] == "Heading 2")
        return paragraphs[start:end]

    assert body("Зачем писать постмортем") == [("Normal", "Выводы по инцидентам забываются."),
                                               ("Normal", "Документ пишется для дежурных.")]
    assert body("Структура") == [("Normal", "Пять частей."), ("List Bullet", "Хронология"),
                                 ("List Bullet", "Причины")]
    assert body("Правила") == [("List Bullet", "Владелец и срок у каждой задачи")]
    assert body("Код") == [("Normal", "Модуль __init__.py и snake_case_name")]

    headings = [text for style, text in paragraphs if style == "Heading 2"]
    # Sections sit where the type's details go; empty common sections are stripped.
    assert headings == ["Кратко", "Участники", "Ключевые цифры", "Зачем писать постмортем",
                        "Структура", "Правила", "Код", "Задачи", "Ход обсуждения"]
